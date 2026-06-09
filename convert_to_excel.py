import json
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

# Read the GeoJSON file
with open('communes.geojson', 'r', encoding='utf-8') as f:
    geojson = json.load(f)

# Create a new workbook
wb = Workbook()
ws = wb.active
ws.title = "Communes"

# Define headers
headers = ['ISO', 'NAME_1 (Arabic)', 'NAME_2 (French)', 'Marocains', 'Etrangers', 'Population', 'Menages']

# Style for header
header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
header_alignment = Alignment(horizontal="center", vertical="center")

# Write headers
for col, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment

# Write data
for row_idx, feature in enumerate(geojson['features'], 2):
    props = feature['properties']
    ws.cell(row=row_idx, column=1, value=props.get('ISO', ''))
    ws.cell(row=row_idx, column=2, value=props.get('NAME_1', ''))
    ws.cell(row=row_idx, column=3, value=props.get('NAME_2', ''))
    ws.cell(row=row_idx, column=4, value=props.get('Marocains_', ''))
    ws.cell(row=row_idx, column=5, value=props.get('Etrangers_', ''))
    ws.cell(row=row_idx, column=6, value=props.get('Populati_1', ''))
    ws.cell(row=row_idx, column=7, value=props.get('Menages_', ''))

# Adjust column widths
ws.column_dimensions['A'].width = 20
ws.column_dimensions['B'].width = 25
ws.column_dimensions['C'].width = 25
ws.column_dimensions['D'].width = 12
ws.column_dimensions['E'].width = 12
ws.column_dimensions['F'].width = 12
ws.column_dimensions['G'].width = 12

# Save the workbook
wb.save('communes.xlsx')
print(f"Successfully converted communes.geojson to communes.xlsx")
print(f"Total features: {len(geojson['features'])}")
